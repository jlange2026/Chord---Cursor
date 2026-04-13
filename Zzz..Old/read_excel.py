#!/usr/bin/env python3
"""Read Excel file using openpyxl"""

import sys

try:
    from openpyxl import load_workbook
    
    # Load the workbook
    wb = load_workbook('Data_2026-01-21-dental.prod.xlsx', data_only=True)
    ws = wb.active
    
    # Get all data
    data = []
    for row in ws.iter_rows(values_only=True):
        data.append(row)
    
    # Print header
    if data:
        headers = data[0]
        print("Columns:", list(headers))
        print(f"\nTotal rows (including header): {len(data)}")
        print("\nFirst 10 rows:")
        for i, row in enumerate(data[:10]):
            print(f"Row {i}: {row}")
    
    # Save as CSV
    import csv
    with open('Data_2026-01-21-dental.prod.csv', 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerows(data)
    print("\n✓ Successfully converted to CSV!")
    
except ImportError:
    print("openpyxl not available. Trying alternative method...")
    
    # Try using zipfile to extract XML
    import zipfile
    import xml.etree.ElementTree as ET
    
    try:
        with zipfile.ZipFile('Data_2026-01-21-dental.prod.xlsx', 'r') as zip_ref:
            # List contents
            print("Excel file structure:")
            for name in zip_ref.namelist()[:20]:
                print(f"  - {name}")
            
            # Try to read shared strings
            try:
                shared_strings_xml = zip_ref.read('xl/sharedStrings.xml')
                print("\n✓ Found shared strings")
            except:
                pass
                
            # Try to read worksheet
            try:
                sheet_xml = zip_ref.read('xl/worksheets/sheet1.xml')
                print("✓ Found worksheet data")
            except:
                pass
                
    except Exception as e:
        print(f"Error reading Excel structure: {e}")
        
except Exception as e:
    print(f"Error: {e}")
    import traceback
    traceback.print_exc()
